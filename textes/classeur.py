import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
d = json.load(open("textes.json"))
NOMS = {"index.html":"Accueil","about-us.html":"À propos","projects.html":"Projets","projects-detail.html":"Détail projet",
        "blog.html":"Blog","blog-detail.html":"Article de blog","contact.html":"Contact","sign-in.html":"Connexion",
        "sign-up.html":"Inscription","privacy-policy.html":"Confidentialité","terms-and-conditions.html":"Conditions","404.html":"Page 404"}
TYPE = {"title":"Titre de l'onglet","description":"Description Google","h1":"Grand titre","h2":"Titre de section","h3":"Sous-titre",
        "h4":"Sous-titre","h5":"Petit titre","h6":"Petit titre","p":"Paragraphe","a":"Lien","button":"Bouton","span":"Libellé",
        "li":"Élément de liste","label":"Étiquette de champ","blockquote":"Citation","textarea[placeholder]":"Texte d'exemple du champ",
        "input[placeholder]":"Texte d'exemple du champ","td":"Cellule","th":"Cellule"}
F = "Arial"; head = Font(name=F, bold=True, color="FFFFFF"); body = Font(name=F); grey = Font(name=F, color="808080")
HF = PatternFill("solid", fgColor="1F2A2E"); YF = PatternFill("solid", fgColor="FFF7C2")
thin = Border(bottom=Side(style="thin", color="E4E5E6"))
wb = Workbook(); ws = wb.active; ws.title = "Mode d'emploi"
for r in [["Comment remplir ce fichier"],[],
          ["1. Chaque onglet correspond à une page du site. L'onglet « Commun » contient le menu et le pied de page, présents sur toutes les pages."],
          ["2. Écris ton texte dans la colonne jaune « Nouveau texte ». Laisse-la vide pour garder le texte d'origine."],
          ["3. Certaines phrases sont découpées en plusieurs lignes (ex. « We create » / « high-performing » / « digital designs… ») : le morceau du milieu est mis en couleur sur le site. Remplis chaque morceau."],
          ["4. Quand c'est fini, renvoie-moi le fichier : je remplace les textes dans le site automatiquement grâce à la colonne « Réf. »."],[],
          ["Exemple", ""],["Texte original","High quality web design solutions you can trust."],["Nouveau texte","Des circuits moto sur mesure, pensés par des passionnés."]]:
    ws.append(r)
ws["A1"].font = Font(name=F, bold=True, size=14)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = body
ws["A8"].font = Font(name=F, bold=True); ws["B10"].fill = YF
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 70
for r in range(3,7): ws.merge_cells(f"A{r}:B{r}"); ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r].height = 32

def sheet(title, rows):
    s = wb.create_sheet(title[:31])
    s.append(["Réf.", "Section", "Type", "Texte original", "Nouveau texte"])
    for c in s[1]: c.font = head; c.fill = HF; c.alignment = Alignment(vertical="center")
    for r in rows:
        s.append(r)
        for i, c in enumerate(s[s.max_row]):
            c.font = grey if i < 3 else body; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = thin
        s.cell(s.max_row, 5).fill = YF
    for col, w in zip("ABCDE", (13, 20, 18, 60, 60)): s.column_dimensions[col].width = w
    s.freeze_panes = "D2"; s.auto_filter.ref = f"A1:E{s.max_row}"
    return len(rows)

def rows_for(page, items, keep):
    out, section, prev, n = [], "Haut de page", "", 0
    for i, (zone, tag, txt) in enumerate(items):
        if zone not in keep or tag.startswith("img"): prev = txt; continue
        if re.fullmatch(r"0\d", prev) and tag == "span": section = txt
        prev = txt
        n += 1
        out.append([f"{page}#{i}", "Menu" if zone=="en-tête" else "Pied de page" if zone=="pied de page" else section,
                    TYPE.get(tag, tag), txt, None])
    return out

total = {}
commun = rows_for("commun", d["index.html"], {"en-tête","pied de page"})
total["Commun"] = sheet("Commun", commun)
for f, nom in NOMS.items():
    total[nom] = sheet(nom, rows_for(f, d[f], {"head","meta","corps"}))
wb.save("Textes du site Studiova.xlsx")
print(total, sum(total.values()))
